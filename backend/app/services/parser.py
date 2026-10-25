"""
Document parsing service.
Extracts raw text from uploaded PDF, DOCX, or TXT files.
"""
import io
import re
import pdfplumber
import docx


class ParsingError(Exception):
    pass


def parse_pdf(file_bytes: bytes) -> str:
    text_parts = []
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
    except Exception as e:
        raise ParsingError(f"Failed to parse PDF: {e}")
    return "\n".join(text_parts)


def parse_docx(file_bytes: bytes) -> str:
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also pull text from tables (common in resumes)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)
        return "\n".join(paragraphs)
    except Exception as e:
        raise ParsingError(f"Failed to parse DOCX: {e}")


def parse_txt(file_bytes: bytes) -> str:
    try:
        return file_bytes.decode("utf-8", errors="ignore")
    except Exception as e:
        raise ParsingError(f"Failed to parse TXT: {e}")


def clean_text(text: str) -> str:
    """Normalize whitespace, remove contact info (URLs, emails, phone numbers), collapse blank lines."""
    text = text.replace("\r", "\n")
    # Strip URLs (including LinkedIn, GitHub, etc.)
    text = re.sub(r"https?://\S+|www\.\S+|linkedin\.com/\S+|github\.com/\S+", "", text, flags=re.IGNORECASE)
    # Strip email addresses
    text = re.sub(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b", "", text)
    # Strip phone numbers (e.g., 9867083433, +1-234-567-8900)
    text = re.sub(r"\b(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}\b", "", text)
    text = re.sub(r"\b\d{10}\b", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_document(filename: str, file_bytes: bytes) -> str:
    """
    Dispatch to the correct parser based on file extension.
    Returns cleaned plain text.
    """
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if ext == "pdf":
        raw = parse_pdf(file_bytes)
    elif ext == "docx":
        raw = parse_docx(file_bytes)
    elif ext == "txt":
        raw = parse_txt(file_bytes)
    else:
        raise ParsingError(
            f"Unsupported file type: .{ext}. Supported types: .pdf, .docx, .txt"
        )

    cleaned = clean_text(raw)
    if not cleaned:
        raise ParsingError(
            "No extractable text found in document (it may be a scanned image)."
        )
    return cleaned
